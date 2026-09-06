
import io
import json
import os
import re
from typing import Dict, List, Tuple

import fitz  # PyMuPDF
import streamlit as st
from docx import Document
from google import genai
from google.genai import types


MODEL_NAME = "gemini-3.5-flash"


def get_api_key() -> str:
    try:
        key = st.secrets["GEMINI_API_KEY"]
    except Exception:
        key = os.getenv("GEMINI_API_KEY")

    if not key:
        raise RuntimeError(
            "GEMINI_API_KEY is not configured. Add it to Streamlit Secrets."
        )
    return str(key)


@st.cache_resource(show_spinner=False)
def get_gemini_client() -> genai.Client:
    return genai.Client(api_key=get_api_key())


def extract_pdf_text(file_bytes: bytes) -> str:
    parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            parts.append(page.get_text("text"))
    return "\n".join(parts).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_txt_text(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace").strip()


def extract_resume_text(uploaded_file) -> str:
    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]
    file_bytes = uploaded_file.getvalue()

    if extension == "pdf":
        return extract_pdf_text(file_bytes)
    if extension == "docx":
        return extract_docx_text(file_bytes)
    if extension == "txt":
        return extract_txt_text(file_bytes)

    raise ValueError("Unsupported file type.")


def normalize_text(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def local_resume_checks(text: str) -> Dict:
    lower = text.lower()
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    section_patterns = {
        "summary": [r"\b(summary|profile|professional summary|objective)\b"],
        "experience": [
            r"\b(experience|work experience|professional experience|employment)\b"
        ],
        "education": [r"\b(education|academic background|qualifications)\b"],
        "skills": [
            r"\b(skills|technical skills|core competencies|competencies)\b"
        ],
        "projects": [r"\b(projects|selected projects|academic projects)\b"],
    }

    sections = {
        name: any(re.search(pattern, lower) for pattern in patterns)
        for name, patterns in section_patterns.items()
    }

    email_present = bool(re.search(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b", text))
    linkedin_present = bool(re.search(r"linkedin\.com", lower))
    github_present = bool(re.search(r"github\.com", lower))
    phone_present = bool(re.search(r"\+?\d[\d\s().-]{7,}\d", text))

    bullet_lines = sum(
        1 for line in lines if re.match(r"^([•●▪◦\-*]|\d+[.)])\s+", line)
    )
    quantified_lines = sum(
        1 for line in lines if re.search(r"\b\d+(?:\.\d+)?%?\b", line)
    )

    action_verbs = [
        "achieved", "analyzed", "built", "created", "developed", "designed",
        "implemented", "improved", "increased", "led", "managed", "optimized",
        "reduced", "streamlined", "launched", "automated", "generated",
        "coordinated", "delivered", "researched", "supported", "trained",
    ]
    action_verb_hits = sum(
        len(re.findall(rf"\b{re.escape(verb)}\b", lower))
        for verb in action_verbs
    )

    return {
        "word_count": len(re.findall(r"\b[\w'-]+\b", text)),
        "line_count": len(lines),
        "sections": sections,
        "email_present": email_present,
        "phone_present": phone_present,
        "linkedin_present": linkedin_present,
        "github_present": github_present,
        "bullet_lines": bullet_lines,
        "quantified_lines": quantified_lines,
        "action_verb_hits": action_verb_hits,
    }


def build_local_score(checks: Dict) -> Tuple[int, List[str]]:
    score = 0
    positives = []
    sections = checks["sections"]

    checks_map = [
        ("email_present", 10, "Email detected"),
        ("phone_present", 5, "Phone number detected"),
        ("linkedin_present", 5, "LinkedIn URL detected"),
        ("github_present", 3, "GitHub URL detected"),
    ]
    for key, points, message in checks_map:
        if checks[key]:
            score += points
            positives.append(message)

    section_points = {
        "experience": (15, "Experience section detected"),
        "education": (10, "Education section detected"),
        "skills": (15, "Skills section detected"),
        "projects": (5, "Projects section detected"),
        "summary": (5, "Summary/Profile section detected"),
    }
    for key, (points, message) in section_points.items():
        if sections[key]:
            score += points
            positives.append(message)

    if checks["bullet_lines"] >= 4:
        score += 10
        positives.append("Bullet-point structure detected")
    elif checks["bullet_lines"] >= 1:
        score += 5
        positives.append("Some bullet points detected")

    if checks["quantified_lines"] >= 3:
        score += 10
        positives.append("Several quantified statements detected")
    elif checks["quantified_lines"] >= 1:
        score += 5
        positives.append("Some quantified content detected")

    if checks["action_verb_hits"] >= 5:
        score += 10
        positives.append("Strong action-verb usage detected")
    elif checks["action_verb_hits"] >= 1:
        score += 5
        positives.append("Some action verbs detected")

    return min(score, 100), positives


def safe_json_loads(text: str) -> Dict:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    return json.loads(cleaned)


def analyze_resume(
    resume_text: str,
    job_description: str,
    client: genai.Client,
) -> Dict:
    local_checks = local_resume_checks(resume_text)
    local_score, _ = build_local_score(local_checks)

    job_text = job_description.strip()
    if not job_text:
        job_text = (
            "No job description was provided. Focus on general ATS compatibility "
            "and clearly state that job-specific keyword matching is limited."
        )

    prompt = f"""
You are a senior recruiter and ATS resume optimization specialist.

Analyze the resume and compare it with the job description.

Return ONLY valid JSON using exactly this structure:
{{
  "ats_score": 0,
  "score_label": "Low",
  "summary": "string",
  "keyword_match": {{
    "matched_keywords": ["string"],
    "missing_keywords": ["string"],
    "notes": "string"
  }},
  "section_feedback": [
    {{
      "section": "string",
      "status": "Good",
      "feedback": "string"
    }}
  ],
  "improvements": [
    {{
      "priority": "High",
      "issue": "string",
      "recommendation": "string"
    }}
  ],
  "rewrites": [
    {{
      "original": "string",
      "improved": "string"
    }}
  ],
  "strengths": ["string"],
  "ats_tips": ["string"]
}}

Rules:
- ats_score must be an integer from 0 to 100.
- This is an ESTIMATED ATS COMPATIBILITY score, not a score from any specific commercial ATS.
- Evaluate keyword relevance, section clarity, parsing-friendly formatting signals,
  measurable achievements, role alignment, and missing information.
- Do not invent candidate facts, employers, dates, skills, metrics, or achievements.
- Suggestions must preserve truthful information.
- If no job description is provided, make that limitation explicit.
- For rewrites, improve existing resume wording only; never fabricate facts.
- Prefer concrete, actionable improvements.

Local parser checks:
{json.dumps(local_checks, indent=2)}

Local heuristic baseline:
{local_score}/100

JOB DESCRIPTION:
{job_text}

RESUME:
{resume_text[:50000]}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json"
        ),
    )

    response_text = response.text
    if not response_text:
        raise RuntimeError("Gemini returned an empty response.")

    result = safe_json_loads(response_text)

    result["ats_score"] = max(
        0, min(100, int(result.get("ats_score", local_score)))
    )
    result.setdefault("score_label", "Moderate")
    result.setdefault("summary", "")
    result.setdefault(
        "keyword_match",
        {"matched_keywords": [], "missing_keywords": [], "notes": ""},
    )
    result.setdefault("section_feedback", [])
    result.setdefault("improvements", [])
    result.setdefault("rewrites", [])
    result.setdefault("strengths", [])
    result.setdefault("ats_tips", [])
    result["local_checks"] = local_checks
    result["local_heuristic_score"] = local_score

    return result


def render_downloadable_report(result: Dict, filename: str) -> None:
    report = {
        "resume_file": filename,
        "estimated_ats_score": result["ats_score"],
        "score_label": result["score_label"],
        "summary": result["summary"],
        "keyword_match": result["keyword_match"],
        "section_feedback": result["section_feedback"],
        "improvements": result["improvements"],
        "rewrites": result["rewrites"],
        "strengths": result["strengths"],
        "ats_tips": result["ats_tips"],
    }

    st.download_button(
        "Download analysis (JSON)",
        data=json.dumps(report, indent=2, ensure_ascii=False),
        file_name="resume_ats_analysis.json",
        mime="application/json",
        use_container_width=True,
    )


st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

st.title("📄 Resume ATS Analyzer")
st.caption(
    "Upload a resume to get an estimated ATS compatibility score, "
    "keyword gaps, section feedback, and practical improvements."
)

with st.sidebar:
    st.header("Job Targeting")
    job_description = st.text_area(
        "Paste the job description (recommended)",
        height=260,
        placeholder="Paste the job posting here for keyword matching...",
    )
    st.info(
        "Different ATS products use different rules. This app provides an "
        "AI-assisted estimate rather than an official ATS score."
    )

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
    help="Text-based PDF, DOCX, or TXT. Scanned image-only PDFs may not extract text.",
)

if uploaded_file:
    try:
        resume_text = normalize_text(extract_resume_text(uploaded_file))

        if not resume_text:
            st.error(
                "No readable text was found. Use a text-based PDF or DOCX, "
                "or upload a TXT resume."
            )
            st.stop()

        word_count = len(re.findall(r"\b[\w'-]+\b", resume_text))
        st.success(f"Loaded {uploaded_file.name} — {word_count} words detected.")

        with st.expander("Preview extracted resume text"):
            st.text(resume_text[:12000])

        analyze_clicked = st.button(
            "🔍 Analyze Resume",
            type="primary",
            use_container_width=True,
        )

        if analyze_clicked:
            try:
                client = get_gemini_client()

                with st.spinner("Analyzing with Gemini Flash..."):
                    result = analyze_resume(
                        resume_text=resume_text,
                        job_description=job_description,
                        client=client,
                    )

                st.session_state["analysis"] = result
                st.session_state["resume_name"] = uploaded_file.name

            except json.JSONDecodeError:
                st.error("Gemini returned an invalid JSON response. Please try again.")
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")

    except Exception as exc:
        st.error(f"Could not read the resume: {exc}")


result = st.session_state.get("analysis")

if result:
    st.divider()

    score = result["ats_score"]
    col1, col2, col3 = st.columns([1, 1, 2])

    with col1:
        st.metric("Estimated ATS Score", f"{score}/100")

    with col2:
        st.metric("Score Level", result["score_label"])

    with col3:
        st.progress(score / 100)

    st.subheader("Overall Assessment")
    st.write(result["summary"])

    keyword_match = result["keyword_match"]
    k1, k2 = st.columns(2)

    with k1:
        st.markdown("### ✅ Matched Keywords")
        matched = keyword_match.get("matched_keywords", [])
        st.write(", ".join(matched) if matched else "No strong matches identified.")

    with k2:
        st.markdown("### ⚠️ Missing / Useful Keywords")
        missing = keyword_match.get("missing_keywords", [])
        st.write(", ".join(missing) if missing else "No major keyword gaps identified.")

    if keyword_match.get("notes"):
        st.caption(keyword_match["notes"])

    st.subheader("Section Feedback")
    for item in result["section_feedback"]:
        st.markdown(
            f"**{item.get('section', 'Section')} — {item.get('status', '')}**"
        )
        st.write(item.get("feedback", ""))

    st.subheader("Priority Improvements")
    for item in result["improvements"]:
        st.markdown(
            f"**{item.get('priority', 'Medium')} priority — "
            f"{item.get('issue', 'Improvement')}**"
        )
        st.write(item.get("recommendation", ""))

    st.subheader("Bullet Rewrites")
    rewrites = result.get("rewrites", [])
    if rewrites:
        for item in rewrites:
            st.markdown(f"**Original:** {item.get('original', '')}")
            st.markdown(f"**Improved:** {item.get('improved', '')}")
            st.divider()
    else:
        st.write("No specific rewrites suggested.")

    c1, c2 = st.columns(2)

    with c1:
        st.markdown("### Strengths")
        for item in result.get("strengths", []):
            st.markdown(f"- {item}")

    with c2:
        st.markdown("### ATS Tips")
        for item in result.get("ats_tips", []):
            st.markdown(f"- {item}")

    with st.expander("Technical parser checks"):
        st.json(result.get("local_checks", {}))
        st.write(
            f"Local heuristic baseline: "
            f"{result.get('local_heuristic_score', 0)}/100"
        )

    render_downloadable_report(
        result,
        st.session_state.get("resume_name", "resume"),
    )
